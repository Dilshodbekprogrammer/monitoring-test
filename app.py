# -*- coding: utf-8 -*-
"""Maktab monitoring test platformasi (MVP).

Rollar: admin, teacher (o'qituvchi), student (o'quvchi).
Ishga tushirish:  python app.py   ->  http://localhost:5000
"""
import hashlib
import hmac
import io
import json
import os
import random
import re
import secrets
import sqlite3
import time
import urllib.error
import urllib.parse
import urllib.request
from datetime import datetime, timedelta, timezone
from functools import wraps

from flask import (Flask, abort, flash, g, jsonify, redirect, render_template,
                   request, send_file, session, url_for)
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from werkzeug.middleware.proxy_fix import ProxyFix
from werkzeug.security import check_password_hash, generate_password_hash

try:  # rasmlarni tekshirish va kichraytirish uchun (tavsiya etiladi)
    from PIL import Image, ImageOps
except ImportError:  # pragma: no cover
    Image = ImageOps = None

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
INSTANCE = os.path.join(BASE_DIR, "instance")
os.makedirs(INSTANCE, exist_ok=True)
DB_PATH = os.environ.get("DB_PATH", os.path.join(INSTANCE, "monitoring.db"))
UPLOAD_DIR = os.environ.get("UPLOAD_DIR", os.path.join(INSTANCE, "uploads"))
for _d in (os.path.dirname(os.path.abspath(DB_PATH)), UPLOAD_DIR):
    try:
        os.makedirs(_d, exist_ok=True)
    except OSError:
        pass

try:
    from zoneinfo import ZoneInfo
    TZ = ZoneInfo(os.environ.get("APP_TZ", "Asia/Tashkent"))
except Exception:  # pragma: no cover
    TZ = timezone(timedelta(hours=5))

LETTERS = "ABCD"
FMT = "%Y-%m-%d %H:%M:%S"
GRACE_SECONDS = 3

# ---- obuna va to'lov sozlamalari (muhit o'zgaruvchilari orqali) ----
PRICE_MONTH = int(os.environ.get("PRICE_MONTH", "49000"))           # 1 oylik narx, so'm
TRIAL_DAYS = int(os.environ.get("TRIAL_DAYS", "7"))                 # bepul sinov davri
MAX_STUDENTS = int(os.environ.get("MAX_STUDENTS_PER_TEACHER", "600"))
HAMYON_BASE = os.environ.get("HAMYON_BASE", "https://hamyon-api.uz").rstrip("/")
HAMYON_SHOP_ID = os.environ.get("HAMYON_SHOP_ID", "").strip()
HAMYON_SHOP_KEY = os.environ.get("HAMYON_SHOP_KEY", "").strip()
UNLIMITED = "2099-12-31 00:00:00"
MONTH_CHOICES = (1, 3, 6, 12)


def now():
    return datetime.now(TZ).replace(tzinfo=None)


def iso(dt):
    return dt.strftime(FMT) if dt else None


def parse(s):
    return datetime.strptime(s, FMT) if s else None


def parse_local_input(s):
    """<input type=datetime-local> qiymatini bazaga mos formatga o'tkazadi."""
    s = (s or "").strip()
    if not s:
        return None
    s = s.replace("T", " ")
    if len(s) == 16:
        s += ":00"
    try:
        datetime.strptime(s, FMT)
    except ValueError:
        return None
    return s


def _secret_key():
    """Sessiya kaliti. Bir nechta worker bir vaqtda ishga tushsa ham hammasi BIR XIL kalit oladi:
    fayl atomik (O_EXCL) yaratiladi, qolganlar uni o'qib oladi. Eng ishonchlisi — SECRET_KEY muhit o'zgaruvchisi."""
    env = os.environ.get("SECRET_KEY", "").strip()
    if env:
        return env, "env"
    path = os.path.join(INSTANCE, "secret.key")
    for _ in range(60):
        try:
            fd = os.open(path, os.O_CREAT | os.O_EXCL | os.O_WRONLY, 0o600)
        except FileExistsError:
            try:
                with open(path) as f:
                    val = f.read().strip()
                if len(val) >= 32:
                    return val, "file"
            except OSError:
                pass
            time.sleep(0.1)  # yaratgan jarayon hali yozib bo'lmagan
            continue
        except OSError:
            break
        key = secrets.token_hex(32)
        with os.fdopen(fd, "w") as f:
            f.write(key)
        return key, "file"
    print("[!] DIQQAT: sessiya kaliti saqlanmadi. SECRET_KEY muhit o'zgaruvchisini bering, aks holda "
          "ko'p workerda foydalanuvchilar tizimdan chiqib ketaveradi.")
    return secrets.token_hex(32), "memory"


SECRET, SECRET_SOURCE = _secret_key()

app = Flask(__name__)
app.config.update(
    SECRET_KEY=SECRET,
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=os.environ.get("COOKIE_SECURE", "") == "1",
    PERMANENT_SESSION_LIFETIME=timedelta(days=14),
    MAX_CONTENT_LENGTH=32 * 1024 * 1024,
)
# nginx / Apache / hosting proksisi orqasida haqiqiy IP va https ni aniqlash uchun
if os.environ.get("TRUST_PROXY", "1") == "1":
    app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1, x_host=1)

IMG_MAX_BYTES = 6 * 1024 * 1024
IMG_MAX_SIDE = 1600
IMG_FIELDS = ("img", "a_img", "b_img", "c_img", "d_img")
MEDIA_RE = re.compile(r"^[0-9a-f]{32}\.(png|jpg)$")

# --------------------------------------------------------------------------
# Ma'lumotlar bazasi
# --------------------------------------------------------------------------
SCHEMA = """
CREATE TABLE IF NOT EXISTS users(
  id INTEGER PRIMARY KEY AUTOINCREMENT,
  username TEXT UNIQUE NOT NULL,
  password_hash TEXT NOT NULL,
  pw_plain TEXT,
  role TEXT NOT NULL CHECK(role IN ('admin','teacher','student')),
  full_name TEXT NOT NULL,
  class_name TEXT,
  active INTEGER NOT NULL DEFAULT 1,
  created_at TEXT,
  phone TEXT,
  owner_id INTEGER,
  sub_until TEXT,
  plan TEXT
);
CREATE TABLE IF NOT EXISTS payments(
  id INTEGER PRIMARY KEY AUTOINCREMENT,
  teacher_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
  order_id TEXT UNIQUE NOT NULL,
  payment_id TEXT,
  months INTEGER NOT NULL,
  amount INTEGER NOT NULL,
  card TEXT,
  status TEXT NOT NULL DEFAULT 'pending',
  created_at TEXT,
  expire_at TEXT,
  prepared_at TEXT,
  paid_at TEXT,
  raw TEXT
);
CREATE TABLE IF NOT EXISTS hook_log(
  id INTEGER PRIMARY KEY AUTOINCREMENT,
  at TEXT,
  path TEXT,
  ip TEXT,
  body TEXT,
  result TEXT
);
CREATE INDEX IF NOT EXISTS idx_pay_teacher ON payments(teacher_id, status);
CREATE INDEX IF NOT EXISTS idx_pay_pid ON payments(payment_id);
CREATE TABLE IF NOT EXISTS subjects(
  id INTEGER PRIMARY KEY AUTOINCREMENT,
  name TEXT UNIQUE NOT NULL
);
CREATE TABLE IF NOT EXISTS teacher_subjects(
  teacher_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
  subject_id INTEGER NOT NULL REFERENCES subjects(id) ON DELETE CASCADE,
  PRIMARY KEY(teacher_id, subject_id)
);
CREATE TABLE IF NOT EXISTS questions(
  id INTEGER PRIMARY KEY AUTOINCREMENT,
  teacher_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
  subject_id INTEGER NOT NULL REFERENCES subjects(id) ON DELETE CASCADE,
  text TEXT NOT NULL,
  a TEXT, b TEXT, c TEXT, d TEXT,
  correct TEXT NOT NULL,
  deleted INTEGER NOT NULL DEFAULT 0,
  created_at TEXT,
  img TEXT, a_img TEXT, b_img TEXT, c_img TEXT, d_img TEXT
);
CREATE TABLE IF NOT EXISTS tests(
  id INTEGER PRIMARY KEY AUTOINCREMENT,
  teacher_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
  subject_id INTEGER NOT NULL REFERENCES subjects(id) ON DELETE CASCADE,
  title TEXT NOT NULL,
  classes TEXT NOT NULL,
  num_questions INTEGER NOT NULL,
  minutes INTEGER NOT NULL,
  start_at TEXT,
  end_at TEXT,
  attempts_allowed INTEGER NOT NULL DEFAULT 1,
  max_warnings INTEGER NOT NULL DEFAULT 3,
  auto_finish INTEGER NOT NULL DEFAULT 1,
  shuffle_q INTEGER NOT NULL DEFAULT 1,
  shuffle_opt INTEGER NOT NULL DEFAULT 1,
  show_result INTEGER NOT NULL DEFAULT 1,
  active INTEGER NOT NULL DEFAULT 1,
  created_at TEXT
);
CREATE TABLE IF NOT EXISTS attempts(
  id INTEGER PRIMARY KEY AUTOINCREMENT,
  test_id INTEGER NOT NULL REFERENCES tests(id) ON DELETE CASCADE,
  student_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
  started_at TEXT NOT NULL,
  finished_at TEXT,
  status TEXT,
  score INTEGER,
  total INTEGER,
  warnings INTEGER NOT NULL DEFAULT 0,
  layout TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS answers(
  attempt_id INTEGER NOT NULL REFERENCES attempts(id) ON DELETE CASCADE,
  question_id INTEGER NOT NULL,
  choice TEXT,
  PRIMARY KEY(attempt_id, question_id)
);
CREATE TABLE IF NOT EXISTS violations(
  id INTEGER PRIMARY KEY AUTOINCREMENT,
  attempt_id INTEGER NOT NULL REFERENCES attempts(id) ON DELETE CASCADE,
  at TEXT,
  kind TEXT
);
CREATE INDEX IF NOT EXISTS idx_q_ts ON questions(teacher_id, subject_id, deleted);
CREATE INDEX IF NOT EXISTS idx_att_test ON attempts(test_id);
CREATE INDEX IF NOT EXISTS idx_att_student ON attempts(student_id);
"""


def db():
    if "db" not in g:
        conn = sqlite3.connect(DB_PATH, timeout=15)
        conn.row_factory = sqlite3.Row
        conn.execute("PRAGMA foreign_keys=ON")
        conn.execute("PRAGMA journal_mode=WAL")
        g.db = conn
    return g.db


@app.teardown_appcontext
def close_db(_exc):
    conn = g.pop("db", None)
    if conn is not None:
        conn.close()


def q(sql, args=(), one=False):
    rows = db().execute(sql, args).fetchall()
    if one:
        return rows[0] if rows else None
    return rows


def ex(sql, args=()):
    cur = db().execute(sql, args)
    db().commit()
    return cur.lastrowid


def _add_column(conn, table, name, typ):
    cols = {r[1] for r in conn.execute("PRAGMA table_info(%s)" % table)}
    if name not in cols:
        try:
            conn.execute("ALTER TABLE %s ADD COLUMN %s %s" % (table, name, typ))
        except sqlite3.OperationalError:
            pass  # boshqa worker allaqachon qo'shgan


def _init_db_once():
    conn = sqlite3.connect(DB_PATH, timeout=30)
    try:
        conn.executescript(SCHEMA)
        # eski bazalarni yangilash
        for name, typ in (("phone", "TEXT"), ("owner_id", "INTEGER"), ("sub_until", "TEXT"), ("plan", "TEXT")):
            _add_column(conn, "users", name, typ)
        for name in IMG_FIELDS:
            _add_column(conn, "questions", name, "TEXT")
        conn.execute("UPDATE users SET sub_until=?, plan='admin' WHERE role='teacher' AND sub_until IS NULL", (UNLIMITED,))
        conn.execute("UPDATE users SET owner_id=(SELECT MIN(id) FROM users WHERE role='teacher') "
                     "WHERE role='student' AND owner_id IS NULL")
        conn.execute("CREATE UNIQUE INDEX IF NOT EXISTS idx_users_phone ON users(phone) WHERE phone IS NOT NULL")
        row = conn.execute("SELECT id FROM users WHERE role='admin' LIMIT 1").fetchone()
        if not row:
            pw = os.environ.get("ADMIN_PASSWORD", "admin123")
            cur = conn.execute(  # OR IGNORE: bir nechta worker bir vaqtda urinsa, xato bermasin
                "INSERT OR IGNORE INTO users(username,password_hash,role,full_name,created_at) VALUES(?,?,?,?,?)",
                ("admin", generate_password_hash(pw), "admin", "Administrator", iso(now())),
            )
            if cur.rowcount:
                print("[!] Admin yaratildi: login=admin, parol=%s  (kirgach almashtiring)" % pw)
        conn.commit()
    finally:
        conn.close()


def init_db():
    for attempt in range(5):
        try:
            return _init_db_once()
        except sqlite3.OperationalError as e:  # "database is locked" — boshqa worker ishlayapti
            if attempt == 4:
                raise
            time.sleep(0.5 * (attempt + 1))


# --------------------------------------------------------------------------
# Xavfsizlik: CSRF, login, rollar
# --------------------------------------------------------------------------
@app.before_request
def csrf_protect():
    if request.method == "POST" and not request.path.startswith("/hooks/"):  # /hooks/ — imzo bilan tekshiriladi
        tok = session.get("csrf")
        sent = request.form.get("csrf") or request.headers.get("X-CSRF")
        if not tok or not sent or not secrets.compare_digest(tok, sent):
            if request.path.startswith("/api/"):
                return jsonify(error="csrf"), 400
            abort(400, "Sessiya muddati tugagan. Sahifani yangilab qayta urinib ko'ring.")


@app.context_processor
def inject_globals():
    if "csrf" not in session:
        session["csrf"] = secrets.token_hex(16)
    u = current_user()
    sub = teacher_sub(u) if u and u["role"] == "teacher" else None
    return dict(csrf_token=session["csrf"], user=u, sub=sub, price_month=PRICE_MONTH)


@app.template_filter("money")
def fmt_money(n):
    try:
        return "{:,}".format(int(n)).replace(",", " ")
    except (TypeError, ValueError):
        return n


@app.template_filter("dt")
def fmt_dt(s):
    if not s:
        return "—"
    try:
        return parse(s).strftime("%d.%m.%Y %H:%M")
    except Exception:
        return s


@app.template_filter("dtinput")
def fmt_dt_input(s):
    return s[:16].replace(" ", "T") if s else ""


def current_user():
    if "user_obj" in g:
        return g.user_obj
    u = None
    if "uid" in session:
        u = q("SELECT * FROM users WHERE id=? AND active=1", (session["uid"],), one=True)
    g.user_obj = u
    return u


def role_required(*roles):
    def deco(f):
        @wraps(f)
        def wrapper(*a, **k):
            u = current_user()
            if not u:
                session.pop("uid", None)
                if request.path.startswith("/api/"):
                    return jsonify(error="auth"), 401
                return redirect(url_for("login", next=request.path))
            if roles and u["role"] not in roles:
                abort(403)
            return f(*a, **k)
        return wrapper
    return deco


FAILS = {}  # (ip, username) -> [urinishlar, bloklangan_vaqt]


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        username = request.form.get("username", "").strip().lower()
        password = request.form.get("password", "")
        key = (request.remote_addr, username)
        rec = FAILS.get(key)
        if rec and rec[1] and rec[1] > now():
            flash("Juda ko'p noto'g'ri urinish. Bir necha daqiqadan keyin qayta urinib ko'ring.", "err")
            return render_template("login.html")
        u = q("SELECT * FROM users WHERE username=? AND active=1", (username,), one=True)
        if u and check_password_hash(u["password_hash"], password):
            FAILS.pop(key, None)
            session.clear()
            session.permanent = True
            session["uid"] = u["id"]
            session["csrf"] = secrets.token_hex(16)
            nxt = request.args.get("next", "")
            if nxt.startswith("/") and not nxt.startswith("//"):
                return redirect(nxt)
            return redirect(url_for("home"))
        rec = rec or [0, None]
        rec[0] += 1
        if rec[0] >= 6:
            rec = [0, now() + timedelta(minutes=5)]
        FAILS[key] = rec
        flash("Login yoki parol noto'g'ri.", "err")
    return render_template("login.html")


@app.route("/logout", methods=["POST"])
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.route("/")
def home():
    u = current_user()
    if not u:
        return render_template("landing.html", trial=TRIAL_DAYS)
    return redirect(url_for({"admin": "admin_home", "teacher": "t_home", "student": "s_home"}[u["role"]]))


@app.errorhandler(400)
@app.errorhandler(403)
@app.errorhandler(404)
@app.errorhandler(413)
@app.errorhandler(500)
def err(e):
    return render_template("error.html", e=e), e.code


# --------------------------------------------------------------------------
# Yordamchi funksiyalar
# --------------------------------------------------------------------------
def grade(pct):
    """5 ballik baho (86+ → 5, 71+ → 4, 55+ → 3, qolgani → 2)."""
    return 5 if pct >= 86 else 4 if pct >= 71 else 3 if pct >= 55 else 2


STATUS_TXT = {
    "finished": "Yakunlangan",
    "time": "Vaqt tugadi",
    "warnings": "Ogohlantirish tufayli yakunlandi",
    None: "Jarayonda",
}

CYR = dict(zip("абвгдезийклмнопрстуфхцэыъь", "abvgdezijklmnoprstufxcey''"))
CYR.update({"ё": "yo", "ж": "j", "ч": "ch", "ш": "sh", "щ": "sh", "ю": "yu", "я": "ya",
            "ў": "o", "қ": "q", "ғ": "g", "ҳ": "h"})


def slugify(s):
    out = []
    for ch in s.lower():
        out.append(CYR.get(ch, ch))
    s = "".join(out)
    return re.sub(r"[^a-z0-9]", "", s)


def make_username(full_name):
    parts = [slugify(p) for p in full_name.split() if slugify(p)]
    base = ".".join(parts[:2]) or "oquvchi"
    name, n = base, 1
    while q("SELECT 1 FROM users WHERE username=?", (name,), one=True):
        n += 1
        name = "%s%d" % (base, n)
    return name


def make_password(n=6):
    alphabet = "abcdefghjkmnpqrstuvwxyz23456789"
    return "".join(secrets.choice(alphabet) for _ in range(n))


def as_int(v, default, lo, hi):
    try:
        x = int(v)
    except (TypeError, ValueError):
        return default
    return max(lo, min(hi, x))


# --------------------------------------------------------------------------
# Obuna
# --------------------------------------------------------------------------
def teacher_sub(u):
    """O'qituvchi obunasi holati."""
    until = parse(u["sub_until"]) if u["sub_until"] else None
    n = now()
    active = bool(until and until > n)
    unlimited = bool(until and until.year >= 2099)
    days_left = int((until - n).total_seconds() // 86400) + 1 if active else 0
    return dict(active=active, until=u["sub_until"], days_left=days_left, plan=u["plan"], unlimited=unlimited)


def teacher_is_active(tid):
    r = q("SELECT sub_until FROM users WHERE id=? AND role='teacher' AND active=1", (tid,), one=True)
    return bool(r and r["sub_until"] and parse(r["sub_until"]) > now())


def extend_subscription(tid, months, plan="paid"):
    """Obunani `months` oyga uzaytiradi (30 kundan). Muddati tugamagan bo'lsa, ustiga qo'shiladi."""
    r = q("SELECT sub_until FROM users WHERE id=?", (tid,), one=True)
    cur = parse(r["sub_until"]) if r and r["sub_until"] else None
    if cur and cur.year >= 2099:
        return cur
    base = max(now(), cur) if cur else now()
    new = base + timedelta(days=30 * months)
    db().execute("UPDATE users SET sub_until=?, plan=? WHERE id=?", (iso(new), plan, tid))
    return new


def sub_required(f):
    """Obunasi tugagan o'qituvchi yangi savol, o'quvchi yoki test qo'sha olmaydi
    (eski natijalarni ko'rish va Excelga yuklash ochiq qoladi)."""
    @wraps(f)
    def wrapper(*a, **k):
        if not teacher_sub(current_user())["active"]:
            flash("Obuna muddati tugagan. Davom etish uchun obunani yangilang. "
                  "Natijalarni ko'rish va yuklab olish ochiq qoladi.", "err")
            return redirect(url_for("t_billing"))
        return f(*a, **k)
    return wrapper


REG_LOG = {}  # ip -> [vaqtlar]


@app.route("/register", methods=["GET", "POST"])
def register():
    if current_user():
        return redirect(url_for("home"))
    form = request.form
    if request.method == "POST":
        ip = request.remote_addr or "?"
        hits = [t for t in REG_LOG.get(ip, []) if t > now() - timedelta(hours=1)]
        full_name = form.get("full_name", "").strip()
        digits = re.sub(r"\D", "", form.get("phone", ""))
        phone = digits[-9:]
        username = form.get("username", "").strip().lower()
        password = form.get("password", "")
        subjects = [s.strip() for s in re.split(r"[,;\n]", form.get("subjects", "")) if s.strip()][:10]
        err = None
        if len(hits) >= 5:
            err = "Juda ko'p urinish. Keyinroq qayta urinib ko'ring."
        elif len(full_name) < 3:
            err = "F.I.Sh. ni kiriting."
        elif len(digits) < 9:
            err = "Telefon raqamini to'liq kiriting (masalan, 90 123 45 67)."
        elif not re.fullmatch(r"[a-z0-9._-]{3,30}", username):
            err = "Login 3–30 belgi: lotin harflari, raqam, nuqta, tire."
        elif len(password) < 8:
            err = "Parol kamida 8 belgi bo'lsin."
        elif not subjects:
            err = "Kamida bitta fan nomini kiriting."
        elif q("SELECT 1 FROM users WHERE username=?", (username,), one=True):
            err = "Bu login band. Boshqasini tanlang."
        elif q("SELECT 1 FROM users WHERE phone=?", (phone,), one=True):
            err = "Bu telefon raqami bilan allaqachon ro'yxatdan o'tilgan. Kirish oynasidan foydalaning."
        if err:
            flash(err, "err")
        else:
            REG_LOG[ip] = hits + [now()]
            until = iso(now() + timedelta(days=TRIAL_DAYS))
            tid = ex("INSERT INTO users(username,password_hash,role,full_name,phone,sub_until,plan,created_at) "
                     "VALUES(?,?,?,?,?,?,?,?)",
                     (username, generate_password_hash(password), "teacher", full_name, phone, until, "trial", iso(now())))
            for name in subjects:
                link_subject(tid, name)
            session.clear()
            session.permanent = True
            session["uid"] = tid
            session["csrf"] = secrets.token_hex(16)
            flash("Xush kelibsiz! %d kunlik bepul sinov davri boshlandi." % TRIAL_DAYS, "ok")
            return redirect(url_for("t_home"))
    return render_template("register.html", form=form, trial=TRIAL_DAYS)


def link_subject(tid, name):
    name = name.strip()[:60]
    row = q("SELECT id FROM subjects WHERE lower(name)=lower(?)", (name,), one=True)
    sid = row["id"] if row else ex("INSERT INTO subjects(name) VALUES(?)", (name,))
    ex("INSERT OR IGNORE INTO teacher_subjects VALUES(?,?)", (tid, sid))
    return sid


@app.route("/teacher/subjects/add", methods=["POST"])
@role_required("teacher")
def t_subject_add():
    names = [s.strip() for s in re.split(r"[,;\n]", request.form.get("name", "")) if s.strip()][:10]
    if not names:
        flash("Fan nomini kiriting.", "err")
    else:
        for n in names:
            link_subject(current_user()["id"], n)
        flash("Fan qo'shildi.", "ok")
    return redirect(url_for("t_home"))


# --------------------------------------------------------------------------
# ADMIN
# --------------------------------------------------------------------------
@app.route("/admin")
@role_required("admin")
def admin_home():
    subjects = q("SELECT * FROM subjects ORDER BY name")
    teachers = q("SELECT * FROM users WHERE role='teacher' ORDER BY full_name")
    tsub = {}
    for r in q("SELECT ts.teacher_id, s.id, s.name FROM teacher_subjects ts JOIN subjects s ON s.id=ts.subject_id"):
        tsub.setdefault(r["teacher_id"], []).append(r["id"])
    subinfo = {t["id"]: teacher_sub(t) for t in teachers}
    weak = check_password_hash(current_user()["password_hash"], "admin123")
    return render_template("admin.html", subjects=subjects, teachers=teachers, tsub=tsub, subinfo=subinfo,
                           weak_admin=weak)


@app.route("/admin/subject", methods=["POST"])
@role_required("admin")
def admin_subject():
    name = request.form.get("name", "").strip()
    if not name:
        flash("Fan nomini kiriting.", "err")
    elif q("SELECT 1 FROM subjects WHERE lower(name)=lower(?)", (name,), one=True):
        flash("Bunday fan allaqachon bor.", "err")
    else:
        ex("INSERT INTO subjects(name) VALUES(?)", (name,))
        flash("Fan qo'shildi.", "ok")
    return redirect(url_for("admin_home"))


@app.route("/admin/teacher", methods=["POST"])
@role_required("admin")
def admin_teacher():
    full_name = request.form.get("full_name", "").strip()
    username = request.form.get("username", "").strip().lower()
    password = request.form.get("password", "")
    sids = request.form.getlist("subject_ids")
    if not full_name or not username or len(password) < 6:
        flash("F.I.Sh., login va kamida 6 belgili parol kiriting.", "err")
    elif q("SELECT 1 FROM users WHERE username=?", (username,), one=True):
        flash("Bu login band.", "err")
    else:
        months = as_int(request.form.get("months"), 0, 0, 120)
        until = iso(now() + timedelta(days=30 * months)) if months else UNLIMITED
        tid = ex(
            "INSERT INTO users(username,password_hash,role,full_name,sub_until,plan,created_at) VALUES(?,?,?,?,?,?,?)",
            (username, generate_password_hash(password), "teacher", full_name, until,
             "paid" if months else "admin", iso(now())),
        )
        for s in sids:
            ex("INSERT OR IGNORE INTO teacher_subjects VALUES(?,?)", (tid, int(s)))
        flash("O'qituvchi qo'shildi.", "ok")
    return redirect(url_for("admin_home"))


@app.route("/admin/teacher/<int:tid>/grant", methods=["POST"])
@role_required("admin")
def admin_teacher_grant(tid):
    if not q("SELECT 1 FROM users WHERE id=? AND role='teacher'", (tid,), one=True):
        abort(404)
    if request.form.get("unlimited"):
        ex("UPDATE users SET sub_until=?, plan='admin' WHERE id=?", (UNLIMITED, tid))
        flash("Cheksiz obuna berildi.", "ok")
    else:
        months = as_int(request.form.get("months"), 1, 1, 120)
        extend_subscription(tid, months)
        db().commit()
        flash("Obuna %d oyga uzaytirildi." % months, "ok")
    return redirect(url_for("admin_home"))


@app.route("/admin/payments")
@role_required("admin")
def admin_payments():
    pays = q("SELECT p.*, u.full_name, u.username FROM payments p JOIN users u ON u.id=p.teacher_id "
             "ORDER BY p.id DESC LIMIT 200")
    logs = q("SELECT * FROM hook_log ORDER BY id DESC LIMIT 30")
    total = q("SELECT COALESCE(SUM(amount),0) s FROM payments WHERE status='paid'", one=True)["s"]
    return render_template("admin_payments.html", pays=pays, logs=logs, total=total)


@app.route("/admin/teacher/<int:tid>/subjects", methods=["POST"])
@role_required("admin")
def admin_teacher_subjects(tid):
    if not q("SELECT 1 FROM users WHERE id=? AND role='teacher'", (tid,), one=True):
        abort(404)
    ex("DELETE FROM teacher_subjects WHERE teacher_id=?", (tid,))
    for s in request.form.getlist("subject_ids"):
        ex("INSERT OR IGNORE INTO teacher_subjects VALUES(?,?)", (tid, int(s)))
    flash("Fanlar yangilandi.", "ok")
    return redirect(url_for("admin_home"))


@app.route("/admin/teacher/<int:tid>/password", methods=["POST"])
@role_required("admin")
def admin_teacher_password(tid):
    pw = request.form.get("password", "")
    if len(pw) < 6:
        flash("Parol kamida 6 belgi bo'lsin.", "err")
    else:
        ex("UPDATE users SET password_hash=? WHERE id=? AND role='teacher'", (generate_password_hash(pw), tid))
        flash("Parol almashtirildi.", "ok")
    return redirect(url_for("admin_home"))


@app.route("/admin/teacher/<int:tid>/toggle", methods=["POST"])
@role_required("admin")
def admin_teacher_toggle(tid):
    ex("UPDATE users SET active=1-active WHERE id=? AND role='teacher'", (tid,))
    return redirect(url_for("admin_home"))


@app.route("/password", methods=["POST"])
@role_required("admin", "teacher")
def change_own_password():
    old, new = request.form.get("old", ""), request.form.get("new", "")
    u = current_user()
    if not check_password_hash(u["password_hash"], old):
        flash("Joriy parol noto'g'ri.", "err")
    elif len(new) < 6:
        flash("Yangi parol kamida 6 belgi bo'lsin.", "err")
    else:
        ex("UPDATE users SET password_hash=? WHERE id=?", (generate_password_hash(new), u["id"]))
        flash("Parol almashtirildi.", "ok")
    return redirect(request.referrer or url_for("home"))


# --------------------------------------------------------------------------
# O'QITUVCHI: savollar
# --------------------------------------------------------------------------
def my_subjects():
    return q(
        "SELECT s.* FROM subjects s JOIN teacher_subjects ts ON ts.subject_id=s.id "
        "WHERE ts.teacher_id=? ORDER BY s.name", (current_user()["id"],))


def own_subject(sid):
    r = q("SELECT s.* FROM subjects s JOIN teacher_subjects ts ON ts.subject_id=s.id "
          "WHERE ts.teacher_id=? AND s.id=?", (current_user()["id"], sid), one=True)
    if not r:
        abort(404)
    return r


def own_question(qid):
    r = q("SELECT * FROM questions WHERE id=? AND teacher_id=? AND deleted=0", (qid, current_user()["id"]), one=True)
    if not r:
        abort(404)
    return r


def own_test(tid):
    r = q("SELECT t.*, s.name AS subject FROM tests t JOIN subjects s ON s.id=t.subject_id "
          "WHERE t.id=? AND t.teacher_id=?", (tid, current_user()["id"]), one=True)
    if not r:
        abort(404)
    return r


@app.route("/teacher")
@role_required("teacher")
def t_home():
    uid = current_user()["id"]
    subjects = q(
        "SELECT s.*, (SELECT COUNT(*) FROM questions qq WHERE qq.subject_id=s.id AND qq.teacher_id=? AND qq.deleted=0) AS qcount "
        "FROM subjects s JOIN teacher_subjects ts ON ts.subject_id=s.id WHERE ts.teacher_id=? ORDER BY s.name",
        (uid, uid))
    tests = q(
        "SELECT t.*, s.name AS subject, "
        "(SELECT COUNT(*) FROM attempts a WHERE a.test_id=t.id AND a.finished_at IS NOT NULL) AS done "
        "FROM tests t JOIN subjects s ON s.id=t.subject_id WHERE t.teacher_id=? ORDER BY t.id DESC LIMIT 5", (uid,))
    nstud = q("SELECT COUNT(*) c FROM users WHERE role='student' AND active=1 AND owner_id=?", (uid,), one=True)["c"]
    return render_template("t_home.html", subjects=subjects, tests=tests, nstud=nstud)


@app.route("/teacher/subject/<int:sid>/questions")
@role_required("teacher")
def t_questions(sid):
    subject = own_subject(sid)
    rows = q("SELECT * FROM questions WHERE teacher_id=? AND subject_id=? AND deleted=0 ORDER BY id DESC",
             (current_user()["id"], sid))
    return render_template("t_questions.html", subject=subject, rows=rows)


class ImgError(ValueError):
    pass


def save_image(fs):
    """Yuklangan JPG/PNG ni tekshiradi, kerak bo'lsa kichraytiradi va saqlaydi. Fayl nomini qaytaradi."""
    raw = fs.read(IMG_MAX_BYTES + 1)
    if not raw:
        raise ImgError("Bo'sh fayl yuklandi.")
    if len(raw) > IMG_MAX_BYTES:
        raise ImgError("Rasm hajmi %d MB dan oshmasin." % (IMG_MAX_BYTES // 1048576))
    if raw.startswith(b"\x89PNG\r\n\x1a\n"):
        ext = "png"
    elif raw.startswith(b"\xff\xd8\xff"):
        ext = "jpg"
    else:
        raise ImgError("Faqat JPG yoki PNG rasm yuklash mumkin.")
    if Image is not None:
        try:
            im = Image.open(io.BytesIO(raw))
            if im.format not in ("PNG", "JPEG"):
                raise ImgError("Faqat JPG yoki PNG rasm yuklash mumkin.")
            if im.size[0] * im.size[1] > 40_000_000:
                raise ImgError("Rasm o'lchami juda katta.")
            im.load()
            im = ImageOps.exif_transpose(im)
            if max(im.size) > IMG_MAX_SIDE:
                im.thumbnail((IMG_MAX_SIDE, IMG_MAX_SIDE))
            out = io.BytesIO()
            if ext == "png":
                if im.mode not in ("RGB", "RGBA", "L", "LA", "P"):
                    im = im.convert("RGBA")
                im.save(out, "PNG", optimize=True)
            else:
                im.convert("RGB").save(out, "JPEG", quality=85, optimize=True)
            raw = out.getvalue()
        except ImgError:
            raise
        except Exception:
            raise ImgError("Rasm fayli buzilgan yoki o'qib bo'lmadi.")
    name = "%s.%s" % (secrets.token_hex(16), ext)
    with open(os.path.join(UPLOAD_DIR, name), "wb") as f:
        f.write(raw)
    return name


def delete_media(name):
    if name and MEDIA_RE.match(name):
        try:
            os.remove(os.path.join(UPLOAD_DIR, name))
        except OSError:
            pass


def question_from_request(old=None):
    """Savol formasini o'qiydi va tekshiradi. (data, err). Variant — matn YOKI rasm bo'lishi mumkin."""
    f = request.form
    data = {k: f.get(k, "").strip() for k in ("text", "a", "b", "c", "d")}
    data["correct"] = f.get("correct", "").strip().upper()
    has = {}
    for fld in IMG_FIELDS:
        up = request.files.get(fld)
        keep = bool(old and old[fld] and not f.get("rm_" + fld))
        has[fld] = bool((up and up.filename) or keep)

    def opt(L):  # L: 'a'..'d'
        return bool(data[L] or has[L + "_img"])

    err = None
    if not data["text"] and not has["img"]:
        err = "Savol matnini kiriting yoki rasm yuklang."
    elif not opt("a") or not opt("b"):
        err = "Kamida A va B variantlarini kiriting (matn yoki rasm)."
    elif data["correct"] not in LETTERS:
        err = "To'g'ri javobni tanlang."
    elif not opt(data["correct"].lower()):
        err = "To'g'ri javob sifatida bo'sh variant tanlangan."
    elif opt("d") and not opt("c"):
        err = "D variantini kiritish uchun avval C ni to'ldiring."
    return data, err


def apply_images(old=None):
    """Yangi rasmlarni saqlaydi. (nomlar, o'chiriladigan_eski_fayllar). Xatoda yangi fayllar o'chiriladi."""
    names, saved = {}, []
    try:
        for fld in IMG_FIELDS:
            up = request.files.get(fld)
            if up and up.filename:
                n = save_image(up)
                saved.append(n)
                names[fld] = n
            elif old and old[fld] and not request.form.get("rm_" + fld):
                names[fld] = old[fld]
            else:
                names[fld] = None
    except ImgError:
        for n in saved:
            delete_media(n)
        raise
    stale = [old[fld] for fld in IMG_FIELDS if old and old[fld] and names[fld] != old[fld]]
    return names, stale


@app.route("/teacher/subject/<int:sid>/questions/new", methods=["GET", "POST"])
@role_required("teacher")
@sub_required
def t_question_new(sid):
    subject = own_subject(sid)
    data = {}
    if request.method == "POST":
        data, err = question_from_request()
        names = {}
        if not err:
            try:
                names, _ = apply_images()
            except ImgError as e:
                err = str(e)
        if err:
            flash(err, "err")
        else:
            ex("INSERT INTO questions(teacher_id,subject_id,text,a,b,c,d,correct,created_at,img,a_img,b_img,c_img,d_img) "
               "VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
               (current_user()["id"], sid, data["text"], data["a"], data["b"], data["c"], data["d"],
                data["correct"], iso(now()), *[names[f] for f in IMG_FIELDS]))
            flash("Savol qo'shildi.", "ok")
            if request.form.get("another"):
                return redirect(url_for("t_question_new", sid=sid))
            return redirect(url_for("t_questions", sid=sid))
    return render_template("t_question_form.html", subject=subject, q=data, edit=False)


@app.route("/teacher/questions/<int:qid>/edit", methods=["GET", "POST"])
@role_required("teacher")
def t_question_edit(qid):
    row = own_question(qid)
    subject = own_subject(row["subject_id"])
    data = dict(row)
    if request.method == "POST":
        data, err = question_from_request(row)
        names, stale = {}, []
        if not err:
            try:
                names, stale = apply_images(row)
            except ImgError as e:
                err = str(e)
        if err:
            flash(err, "err")
            data.update({f: row[f] for f in IMG_FIELDS})  # forma qayta chizilganda eski rasmlar ko'rinsin
        else:
            ex("UPDATE questions SET text=?,a=?,b=?,c=?,d=?,correct=?,img=?,a_img=?,b_img=?,c_img=?,d_img=? WHERE id=?",
               (data["text"], data["a"], data["b"], data["c"], data["d"], data["correct"],
                *[names[f] for f in IMG_FIELDS], qid))
            for n in stale:
                delete_media(n)
            flash("Savol yangilandi.", "ok")
            return redirect(url_for("t_questions", sid=row["subject_id"]))
    return render_template("t_question_form.html", subject=subject, q=data, edit=True)


@app.route("/teacher/questions/<int:qid>/delete", methods=["POST"])
@role_required("teacher")
def t_question_delete(qid):
    row = own_question(qid)
    ex("UPDATE questions SET deleted=1 WHERE id=?", (qid,))
    flash("Savol o'chirildi.", "ok")
    return redirect(url_for("t_questions", sid=row["subject_id"]))


@app.route("/teacher/subject/<int:sid>/questions/clear", methods=["POST"])
@role_required("teacher")
def t_questions_clear(sid):
    own_subject(sid)
    ex("UPDATE questions SET deleted=1 WHERE teacher_id=? AND subject_id=?", (current_user()["id"], sid))
    flash("Barcha savollar o'chirildi.", "ok")
    return redirect(url_for("t_questions", sid=sid))


# ---- hujjatdan savollarni yuklash ----------------------------------------
def _norm_letter(v):
    v = (v or "").strip().upper()
    m = re.match(r"^([A-D])\b", v)
    if m:
        return m.group(1)
    # kirill A,B,V... ga yo'l qo'ymaymiz; faqat lotin
    return v[:1] if v[:1] in LETTERS else ""


def parse_xlsx(fileobj):
    wb = load_workbook(fileobj, data_only=True)
    ws = wb.active
    rows = []
    for i, r in enumerate(ws.iter_rows(min_row=1, values_only=True), start=1):
        cells = ["" if c is None else str(c).strip() for c in (list(r) + [None] * 6)[:6]]
        if i == 1 and _norm_letter(cells[5]) == "":
            continue  # sarlavha qatori
        if not any(cells):
            continue
        rows.append((i, cells))
    return rows


def parse_docx(fileobj):
    import docx
    d = docx.Document(fileobj)
    rows = []
    # 1) jadval: Savol | A | B | C | D | To'g'ri javob
    for t in d.tables:
        for i, r in enumerate(t.rows, start=1):
            cells = [c.text.strip() for c in r.cells][:6]
            cells += [""] * (6 - len(cells))
            if i == 1 and _norm_letter(cells[5]) == "":
                continue
            if any(cells):
                rows.append((len(rows) + 1, cells))
    if rows:
        return rows
    # 2) matn:  1. Savol / A) ... / B) ... / Javob: B
    lines = []
    for p in d.paragraphs:
        for ln in p.text.split("\n"):
            if ln.strip():
                lines.append(ln.strip())
    cur = None
    opt_re = re.compile(r"^\*?\s*([A-Da-d])\s*[\).:]\s*(.+)$")
    ans_re = re.compile(r"^(?:to['’`ʻ]?g['’`ʻ]?ri\s+)?(?:javob|answer|ans)\s*[:\-–]\s*([A-Da-d])\b", re.I)
    q_re = re.compile(r"^\d+\s*[\.\)]\s*(.+)$")
    out = []

    def flush():
        nonlocal cur
        if cur:
            out.append((len(out) + 1, [cur["t"], cur["A"], cur["B"], cur["C"], cur["D"], cur["ans"]]))
        cur = None

    for ln in lines:
        m = ans_re.match(ln)
        if m and cur:
            cur["ans"] = m.group(1).upper()
            continue
        m = opt_re.match(ln)
        if m and cur:
            L = m.group(1).upper()
            cur[L] = m.group(2).strip()
            if ln.lstrip().startswith("*"):
                cur["ans"] = L
            continue
        m = q_re.match(ln)
        if m:
            flush()
            cur = {"t": m.group(1).strip(), "A": "", "B": "", "C": "", "D": "", "ans": ""}
            continue
        if cur and not cur["A"]:
            cur["t"] += " " + ln  # ko'p qatorli savol matni
    flush()
    return out


@app.route("/teacher/subject/<int:sid>/upload", methods=["POST"])
@role_required("teacher")
@sub_required
def t_upload(sid):
    own_subject(sid)
    f = request.files.get("file")
    if not f or not f.filename:
        flash("Fayl tanlanmadi.", "err")
        return redirect(url_for("t_questions", sid=sid))
    ext = f.filename.rsplit(".", 1)[-1].lower()
    try:
        data = io.BytesIO(f.read())
        if ext == "xlsx":
            rows = parse_xlsx(data)
        elif ext == "docx":
            rows = parse_docx(data)
        else:
            flash("Faqat .xlsx yoki .docx fayl yuklang.", "err")
            return redirect(url_for("t_questions", sid=sid))
    except Exception as e:  # noqa
        flash("Faylni o'qib bo'lmadi: %s" % e, "err")
        return redirect(url_for("t_questions", sid=sid))
    added, errors = 0, []
    uid = current_user()["id"]
    for idx, c in rows:
        text, a, b, cc, d, corr = c
        corr = _norm_letter(corr)
        opts = {"A": a, "B": b, "C": cc, "D": d}
        if not text:
            errors.append("%d-qator: savol matni yo'q" % idx)
        elif not a or not b:
            errors.append("%d-qator: kamida A va B kerak" % idx)
        elif corr not in LETTERS or not opts[corr]:
            errors.append("%d-qator: to'g'ri javob noto'g'ri" % idx)
        elif d and not cc:
            errors.append("%d-qator: D bor, C yo'q" % idx)
        else:
            db().execute(
                "INSERT INTO questions(teacher_id,subject_id,text,a,b,c,d,correct,created_at) VALUES(?,?,?,?,?,?,?,?,?)",
                (uid, sid, text, a, b, cc, d, corr, iso(now())))
            added += 1
    db().commit()
    if added:
        flash("%d ta savol yuklandi." % added, "ok")
    if errors:
        flash("O'tkazib yuborilgan qatorlar (%d): %s%s" % (
            len(errors), "; ".join(errors[:8]), " ..." if len(errors) > 8 else ""), "err")
    if not added and not errors:
        flash("Fayldan savol topilmadi. Shablon formatini tekshiring.", "err")
    return redirect(url_for("t_questions", sid=sid))


@app.route("/teacher/template.xlsx")
@role_required("teacher")
def t_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "Savollar"
    ws.append(["Savol", "A", "B", "C", "D", "To'g'ri javob (A/B/C/D)"])
    ws.append(["2 + 2 nechiga teng?", "3", "4", "5", "6", "B"])
    ws.append(["O'zbekiston poytaxti qaysi shahar?", "Samarqand", "Buxoro", "Toshkent", "Namangan", "C"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F5BEA")
    for i, w in enumerate([60, 22, 22, 22, 22, 24], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return send_file(bio, as_attachment=True, download_name="savollar_shablon.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# --------------------------------------------------------------------------
# O'QITUVCHI: o'quvchilar
# --------------------------------------------------------------------------
def class_list():
    return q("SELECT class_name, COUNT(*) AS n FROM users WHERE role='student' AND active=1 AND owner_id=? "
             "GROUP BY class_name ORDER BY class_name", (current_user()["id"],))


def own_student(sid):
    s = q("SELECT * FROM users WHERE id=? AND role='student' AND owner_id=?", (sid, current_user()["id"]), one=True)
    if not s:
        abort(404)
    return s


@app.route("/teacher/students")
@role_required("teacher")
def t_students():
    cls = request.args.get("class", "")
    classes = class_list()
    students = []
    if cls:
        students = q("SELECT * FROM users WHERE role='student' AND active=1 AND owner_id=? AND class_name=? "
                     "ORDER BY full_name", (current_user()["id"], cls))
    return render_template("t_students.html", classes=classes, cls=cls, students=students)


@app.route("/teacher/students/add", methods=["POST"])
@role_required("teacher")
@sub_required
def t_students_add():
    uid = current_user()["id"]
    cls = request.form.get("class_name", "").strip()[:30]
    names = [re.sub(r"^\s*\d+[\.\)]\s*", "", ln).strip()[:80] for ln in request.form.get("names", "").splitlines()]
    names = [n for n in names if n]
    if not cls or not names:
        flash("Sinf nomini va kamida bitta ism-familiya kiriting.", "err")
        return redirect(url_for("t_students", **({"class": cls} if cls else {})))
    have = q("SELECT COUNT(*) c FROM users WHERE role='student' AND active=1 AND owner_id=?", (uid,), one=True)["c"]
    if have + len(names) > MAX_STUDENTS:
        flash("O'quvchilar soni chegarasi: %d ta." % MAX_STUDENTS, "err")
        return redirect(url_for("t_students", **{"class": cls}))
    for n in names:
        pw = make_password()
        ex("INSERT INTO users(username,password_hash,pw_plain,role,full_name,class_name,owner_id,created_at) "
           "VALUES(?,?,?,?,?,?,?,?)",
           (make_username(n), generate_password_hash(pw), pw, "student", n, cls, uid, iso(now())))
    flash("%d ta o'quvchi qo'shildi. Login-parollar quyida." % len(names), "ok")
    return redirect(url_for("t_students", **{"class": cls}))


@app.route("/teacher/students/<int:sid>/reset", methods=["POST"])
@role_required("teacher")
def t_student_reset(sid):
    s = own_student(sid)
    pw = make_password()
    ex("UPDATE users SET password_hash=?, pw_plain=? WHERE id=?", (generate_password_hash(pw), pw, sid))
    flash("%s uchun yangi parol: %s" % (s["full_name"], pw), "ok")
    return redirect(url_for("t_students", **{"class": s["class_name"]}))


@app.route("/teacher/students/<int:sid>/delete", methods=["POST"])
@role_required("teacher")
def t_student_delete(sid):
    s = own_student(sid)
    ex("UPDATE users SET active=0 WHERE id=?", (sid,))
    flash("O'quvchi arxivlandi (natijalari saqlanadi).", "ok")
    return redirect(url_for("t_students", **{"class": s["class_name"]}))


def _autosize(ws, min_w=8, max_w=60):
    for col in ws.columns:
        width = max((len(str(c.value)) if c.value is not None else 0) for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max(min_w, min(max_w, width + 3))


def _style_header(ws, row=1):
    for c in ws[row]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F5BEA")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _send_wb(wb, name):
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return send_file(bio, as_attachment=True, download_name=name,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.route("/teacher/students/credentials.xlsx")
@role_required("teacher")
def t_credentials():
    cls = request.args.get("class", "")
    rows = q("SELECT * FROM users WHERE role='student' AND active=1 AND owner_id=? AND class_name=? ORDER BY full_name",
             (current_user()["id"], cls))
    if not rows:
        abort(404)
    wb = Workbook()
    ws = wb.active
    ws.title = cls[:30] or "Sinf"
    ws.append(["№", "F.I.Sh.", "Sinf", "Login", "Parol"])
    for i, r in enumerate(rows, start=1):
        ws.append([i, r["full_name"], r["class_name"], r["username"], r["pw_plain"] or "(o'zgartirilgan)"])
    _style_header(ws)
    _autosize(ws)
    return _send_wb(wb, "login_parollar_%s.xlsx" % (slugify(cls) or "sinf"))


# --------------------------------------------------------------------------
# O'QITUVCHI: testlar
# --------------------------------------------------------------------------
@app.route("/teacher/tests")
@role_required("teacher")
def t_tests():
    rows = q(
        "SELECT t.*, s.name AS subject, "
        "(SELECT COUNT(*) FROM attempts a WHERE a.test_id=t.id AND a.finished_at IS NOT NULL) AS done, "
        "(SELECT COUNT(*) FROM attempts a WHERE a.test_id=t.id AND a.finished_at IS NULL) AS running "
        "FROM tests t JOIN subjects s ON s.id=t.subject_id WHERE t.teacher_id=? ORDER BY t.id DESC",
        (current_user()["id"],))
    return render_template("t_tests.html", rows=rows)


@app.route("/teacher/tests/new", methods=["GET", "POST"])
@role_required("teacher")
@sub_required
def t_test_new():
    uid = current_user()["id"]
    subjects = my_subjects()
    classes = class_list()
    bank = {r["subject_id"]: r["n"] for r in q(
        "SELECT subject_id, COUNT(*) n FROM questions WHERE teacher_id=? AND deleted=0 GROUP BY subject_id", (uid,))}
    form = {}
    if request.method == "POST":
        f = request.form
        form = f
        title = f.get("title", "").strip()
        sid = as_int(f.get("subject_id"), 0, 0, 10**9)
        sel = [c for c in f.getlist("classes") if c]
        n = as_int(f.get("num_questions"), 0, 0, 500)
        minutes = as_int(f.get("minutes"), 0, 0, 600)
        err = None
        if not title:
            err = "Test nomini kiriting."
        elif not any(s["id"] == sid for s in subjects):
            err = "Fan tanlang."
        elif not sel:
            err = "Kamida bitta sinfni tanlang."
        elif n < 1:
            err = "Savollar soni kamida 1 bo'lsin."
        elif n > bank.get(sid, 0):
            err = "Savollar bankida faqat %d ta savol bor (siz %d ta tanladingiz)." % (bank.get(sid, 0), n)
        elif minutes < 1:
            err = "Vaqtni daqiqada kiriting (kamida 1)."
        start, end = parse_local_input(f.get("start_at")), parse_local_input(f.get("end_at"))
        if not err and start and end and end <= start:
            err = "Tugash vaqti boshlanish vaqtidan keyin bo'lishi kerak."
        if err:
            flash(err, "err")
        else:
            tid = ex(
                "INSERT INTO tests(teacher_id,subject_id,title,classes,num_questions,minutes,start_at,end_at,"
                "attempts_allowed,max_warnings,auto_finish,shuffle_q,shuffle_opt,show_result,active,created_at) "
                "VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,1,?)",
                (uid, sid, title, "|" + "|".join(sel) + "|", n, minutes, start, end,
                 as_int(f.get("attempts_allowed"), 1, 1, 10), as_int(f.get("max_warnings"), 3, 1, 10),
                 1 if f.get("auto_finish") else 0, 1 if f.get("shuffle_q") else 0,
                 1 if f.get("shuffle_opt") else 0, 1 if f.get("show_result") else 0, iso(now())))
            flash("Test yaratildi. O'quvchilar o'z akkauntidan ishlay oladi.", "ok")
            return redirect(url_for("t_tests"))
    return render_template("t_test_form.html", subjects=subjects, classes=classes, bank=bank, form=form,
                           sel=request.form.getlist("classes"))


@app.route("/teacher/tests/<int:tid>/toggle", methods=["POST"])
@role_required("teacher")
def t_test_toggle(tid):
    own_test(tid)
    ex("UPDATE tests SET active=1-active WHERE id=?", (tid,))
    return redirect(url_for("t_tests"))


@app.route("/teacher/tests/<int:tid>/delete", methods=["POST"])
@role_required("teacher")
def t_test_delete(tid):
    own_test(tid)
    ex("DELETE FROM tests WHERE id=?", (tid,))
    flash("Test va uning natijalari o'chirildi.", "ok")
    return redirect(url_for("t_tests"))


def results_for(test):
    for a in q("SELECT id FROM attempts WHERE test_id=? AND finished_at IS NULL", (test["id"],)):
        ensure_current(a["id"])
    rows = q("SELECT a.*, u.full_name, u.class_name FROM attempts a JOIN users u ON u.id=a.student_id "
             "WHERE a.test_id=? ORDER BY u.class_name, u.full_name, a.id", (test["id"],))
    out = []
    for a in rows:
        pct = round(100.0 * a["score"] / a["total"], 1) if a["finished_at"] and a["total"] else None
        dur = None
        if a["finished_at"]:
            dur = round((parse(a["finished_at"]) - parse(a["started_at"])).total_seconds() / 60.0, 1)
        out.append(dict(a=a, pct=pct, grade=grade(pct) if pct is not None else None, dur=dur,
                        status=STATUS_TXT.get(a["status"] if a["finished_at"] else None, "Jarayonda")))
    classes = [c for c in test["classes"].split("|") if c]
    taken = {a["student_id"] for a in rows}
    missing = [s for s in q(
        "SELECT * FROM users WHERE role='student' AND active=1 AND owner_id=? AND class_name IN (%s) "
        "ORDER BY class_name, full_name" % ",".join("?" * len(classes)),
        [test["teacher_id"]] + classes) if s["id"] not in taken] if classes else []
    return out, missing


@app.route("/teacher/tests/<int:tid>/results")
@role_required("teacher")
def t_results(tid):
    test = own_test(tid)
    res, missing = results_for(test)
    fin = [r for r in res if r["pct"] is not None]
    stats = None
    if fin:
        pcts = [r["pct"] for r in fin]
        stats = dict(n=len(fin), avg=round(sum(pcts) / len(pcts), 1), mx=max(pcts), mn=min(pcts),
                     warn=sum(1 for r in fin if r["a"]["warnings"] > 0))
    return render_template("t_results.html", test=test, res=res, missing=missing, stats=stats)


@app.route("/teacher/tests/<int:tid>/results.xlsx")
@role_required("teacher")
def t_results_xlsx(tid):
    test = own_test(tid)
    res, missing = results_for(test)
    wb = Workbook()
    ws = wb.active
    ws.title = "Natijalar"
    ws.append(["Test: %s" % test["title"], "Fan: %s" % test["subject"]])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(["№", "F.I.Sh.", "Sinf", "To'g'ri javoblar", "Savollar soni", "Foiz (%)", "Baho",
               "Ogohlantirishlar", "Boshlangan", "Tugagan", "Davomiyligi (daq)", "Holat"])
    hdr = ws.max_row
    for i, r in enumerate(res, start=1):
        a = r["a"]
        ws.append([i, a["full_name"], a["class_name"], a["score"], a["total"], r["pct"], r["grade"],
                   a["warnings"], fmt_dt(a["started_at"]), fmt_dt(a["finished_at"]), r["dur"], r["status"]])
    _style_header(ws, hdr)
    ws.freeze_panes = ws.cell(row=hdr + 1, column=1)
    _autosize(ws)
    ws2 = wb.create_sheet("Ishtirok etmaganlar")
    ws2.append(["№", "F.I.Sh.", "Sinf"])
    for i, s in enumerate(missing, start=1):
        ws2.append([i, s["full_name"], s["class_name"]])
    _style_header(ws2)
    _autosize(ws2)
    return _send_wb(wb, "natijalar_%s.xlsx" % (slugify(test["title"]) or "test"))


@app.route("/teacher/attempts/<int:aid>")
@role_required("teacher")
def t_attempt(aid):
    att = q("SELECT a.*, u.full_name, u.class_name FROM attempts a JOIN users u ON u.id=a.student_id WHERE a.id=?",
            (aid,), one=True)
    if not att:
        abort(404)
    test = own_test(att["test_id"])
    ensure_current(aid)
    att = q("SELECT a.*, u.full_name, u.class_name FROM attempts a JOIN users u ON u.id=a.student_id WHERE a.id=?",
            (aid,), one=True)
    items = review_items(att)
    viol = q("SELECT * FROM violations WHERE attempt_id=? ORDER BY id", (aid,))
    return render_template("t_attempt.html", test=test, att=att, items=items, viol=viol,
                           status=STATUS_TXT.get(att["status"] if att["finished_at"] else None))


def review_items(att):
    layout = json.loads(att["layout"])
    ans = {r["question_id"]: r["choice"] for r in q("SELECT * FROM answers WHERE attempt_id=?", (att["id"],))}
    items = []
    for it in layout:
        qr = q("SELECT * FROM questions WHERE id=?", (it["qid"],), one=True)
        if not qr:
            continue
        items.append(dict(q=qr, order=it["order"], chosen=ans.get(it["qid"]),
                          ok=ans.get(it["qid"]) == qr["correct"]))
    return items


# --------------------------------------------------------------------------
# O'QUVCHI
# --------------------------------------------------------------------------
def test_deadline(att, test):
    d = parse(att["started_at"]) + timedelta(minutes=test["minutes"])
    if test["end_at"]:
        d = min(d, parse(test["end_at"]))
    return d


def finish_attempt(aid, status):
    att = q("SELECT * FROM attempts WHERE id=?", (aid,), one=True)
    if not att or att["finished_at"]:
        return
    layout = json.loads(att["layout"])
    ans = {r["question_id"]: r["choice"] for r in q("SELECT * FROM answers WHERE attempt_id=?", (aid,))}
    correct = 0
    for it in layout:
        qr = q("SELECT correct FROM questions WHERE id=?", (it["qid"],), one=True)
        if qr and ans.get(it["qid"]) == qr["correct"]:
            correct += 1
    ex("UPDATE attempts SET finished_at=?, status=?, score=?, total=? WHERE id=? AND finished_at IS NULL",
       (iso(now()), status, correct, len(layout), aid))


def ensure_current(aid):
    """Vaqti o'tib ketgan tugallanmagan urinishni yopadi va yangilangan qatorni qaytaradi."""
    att = q("SELECT * FROM attempts WHERE id=?", (aid,), one=True)
    if att and not att["finished_at"]:
        test = q("SELECT * FROM tests WHERE id=?", (att["test_id"],), one=True)
        if now() > test_deadline(att, test) + timedelta(seconds=GRACE_SECONDS):
            finish_attempt(aid, "time")
            att = q("SELECT * FROM attempts WHERE id=?", (aid,), one=True)
    return att


def own_attempt(aid):
    att = q("SELECT * FROM attempts WHERE id=? AND student_id=?", (aid, current_user()["id"]), one=True)
    if not att:
        abort(404)
    return ensure_current(aid)


@app.route("/student")
@role_required("student")
def s_home():
    u = current_user()
    n = now()
    tests = q("SELECT t.*, s.name AS subject, us.full_name AS teacher FROM tests t "
              "JOIN subjects s ON s.id=t.subject_id JOIN users us ON us.id=t.teacher_id "
              "WHERE t.active=1 AND t.teacher_id=? AND t.classes LIKE ? ORDER BY t.id DESC",
              (u["owner_id"], "%|" + (u["class_name"] or "") + "|%"))
    cards = []
    for t in tests:
        atts = [ensure_current(a["id"]) for a in q(
            "SELECT id FROM attempts WHERE test_id=? AND student_id=? ORDER BY id", (t["id"], u["id"]))]
        running = next((a for a in atts if not a["finished_at"]), None)
        done = [a for a in atts if a["finished_at"]]
        start, end = parse(t["start_at"]), parse(t["end_at"])
        if end and n > end and not running:
            state = "closed"
        elif start and n < start:
            state = "soon"
        elif running:
            state = "running"
        elif len(done) >= t["attempts_allowed"]:
            state = "done"
        elif not teacher_is_active(t["teacher_id"]):
            state = "suspended"
        else:
            state = "open"
        cards.append(dict(t=t, state=state, running=running, done=done))
    return render_template("s_home.html", cards=cards)


def build_layout(test):
    rows = q("SELECT * FROM questions WHERE teacher_id=? AND subject_id=? AND deleted=0",
             (test["teacher_id"], test["subject_id"]))
    random.shuffle(rows)
    rows = rows[:test["num_questions"]]
    if not test["shuffle_q"]:
        rows.sort(key=lambda r: r["id"])
    layout = []
    for r in rows:
        letters = [L for L in LETTERS if r[L.lower()] or r[L.lower() + "_img"]]
        if test["shuffle_opt"]:
            random.shuffle(letters)
        layout.append({"qid": r["id"], "order": letters})
    return layout


@app.route("/student/test/<int:tid>/start", methods=["POST"])
@role_required("student")
def s_start(tid):
    u = current_user()
    t = q("SELECT * FROM tests WHERE id=? AND active=1 AND teacher_id=? AND classes LIKE ?",
          (tid, u["owner_id"], "%|" + (u["class_name"] or "") + "|%"), one=True)
    if not t:
        abort(404)
    n = now()
    for a in q("SELECT id FROM attempts WHERE test_id=? AND student_id=? AND finished_at IS NULL", (tid, u["id"])):
        a = ensure_current(a["id"])
        if not a["finished_at"]:
            return redirect(url_for("s_take", aid=a["id"]))
    if not teacher_is_active(t["teacher_id"]):
        flash("Bu test hozircha mavjud emas. O'qituvchingizga murojaat qiling.", "err")
        return redirect(url_for("s_home"))
    if t["start_at"] and n < parse(t["start_at"]):
        flash("Test hali boshlanmagan.", "err")
        return redirect(url_for("s_home"))
    if t["end_at"] and n >= parse(t["end_at"]):
        flash("Test yopilgan.", "err")
        return redirect(url_for("s_home"))
    used = q("SELECT COUNT(*) c FROM attempts WHERE test_id=? AND student_id=?", (tid, u["id"]), one=True)["c"]
    if used >= t["attempts_allowed"]:
        flash("Urinishlar soni tugagan.", "err")
        return redirect(url_for("s_home"))
    layout = build_layout(t)
    if not layout:
        flash("Testda savollar yo'q. O'qituvchiga murojaat qiling.", "err")
        return redirect(url_for("s_home"))
    aid = ex("INSERT INTO attempts(test_id,student_id,started_at,layout,total) VALUES(?,?,?,?,?)",
             (tid, u["id"], iso(n), json.dumps(layout), len(layout)))
    return redirect(url_for("s_take", aid=aid))


@app.route("/student/attempt/<int:aid>")
@role_required("student")
def s_take(aid):
    att = own_attempt(aid)
    if att["finished_at"]:
        return redirect(url_for("s_result", aid=aid))
    test = q("SELECT t.*, s.name AS subject FROM tests t JOIN subjects s ON s.id=t.subject_id WHERE t.id=?",
             (att["test_id"],), one=True)
    layout = json.loads(att["layout"])
    ans = {r["question_id"]: r["choice"] for r in q("SELECT * FROM answers WHERE attempt_id=?", (aid,))}
    items = []
    for it in layout:
        qr = q("SELECT * FROM questions WHERE id=?", (it["qid"],), one=True)
        opts = [(L, qr[L.lower()], qr[L.lower() + "_img"]) for L in it["order"]]
        items.append(dict(id=qr["id"], text=qr["text"], img=qr["img"], opts=opts, chosen=ans.get(qr["id"])))
    remaining = int((test_deadline(att, test) - now()).total_seconds())
    return render_template("s_take.html", att=att, test=test, items=items, remaining=max(remaining, 0))


@app.route("/student/attempt/<int:aid>/result")
@role_required("student")
def s_result(aid):
    att = own_attempt(aid)
    if not att["finished_at"]:
        return redirect(url_for("s_take", aid=aid))
    test = q("SELECT t.*, s.name AS subject FROM tests t JOIN subjects s ON s.id=t.subject_id WHERE t.id=?",
             (att["test_id"],), one=True)
    pct = round(100.0 * att["score"] / att["total"], 1) if att["total"] else 0
    return render_template("s_result.html", att=att, test=test, pct=pct, grade=grade(pct),
                           status=STATUS_TXT.get(att["status"]))


# ---- API (test paytida) --------------------------------------------------
def api_attempt(aid):
    att = q("SELECT * FROM attempts WHERE id=? AND student_id=?", (aid, current_user()["id"]), one=True)
    if not att:
        return None, None
    att = ensure_current(aid)
    test = q("SELECT * FROM tests WHERE id=?", (att["test_id"],), one=True)
    return att, test


@app.route("/api/attempt/<int:aid>/answer", methods=["POST"])
@role_required("student")
def api_answer(aid):
    att, test = api_attempt(aid)
    if not att:
        return jsonify(error="not_found"), 404
    if att["finished_at"]:
        return jsonify(finished=True)
    data = request.get_json(silent=True) or {}
    qid, choice = data.get("qid"), str(data.get("choice", "")).upper()
    layout = json.loads(att["layout"])
    if choice not in LETTERS or not any(it["qid"] == qid and choice in it["order"] for it in layout):
        return jsonify(error="bad"), 400
    ex("INSERT INTO answers(attempt_id,question_id,choice) VALUES(?,?,?) "
       "ON CONFLICT(attempt_id,question_id) DO UPDATE SET choice=excluded.choice", (aid, qid, choice))
    return jsonify(ok=True)


@app.route("/api/attempt/<int:aid>/violation", methods=["POST"])
@role_required("student")
def api_violation(aid):
    att, test = api_attempt(aid)
    if not att:
        return jsonify(error="not_found"), 404
    if att["finished_at"]:
        return jsonify(finished=True, warnings=att["warnings"], max=test["max_warnings"])
    kind = (request.get_json(silent=True) or {}).get("kind", "blur")[:30]
    warnings = att["warnings"] + 1
    ex("UPDATE attempts SET warnings=? WHERE id=?", (warnings, aid))
    ex("INSERT INTO violations(attempt_id,at,kind) VALUES(?,?,?)", (aid, iso(now()), kind))
    finished = False
    if test["auto_finish"] and warnings >= test["max_warnings"]:
        finish_attempt(aid, "warnings")
        finished = True
    return jsonify(finished=finished, warnings=warnings, max=test["max_warnings"])


@app.route("/api/attempt/<int:aid>/finish", methods=["POST"])
@role_required("student")
def api_finish(aid):
    att, test = api_attempt(aid)
    if not att:
        return jsonify(error="not_found"), 404
    finish_attempt(aid, "finished")
    return jsonify(finished=True, url=url_for("s_result", aid=aid))


# --------------------------------------------------------------------------
# TO'LOV (HamyonAPI): obuna sotib olish
# --------------------------------------------------------------------------
class PayError(Exception):
    pass


def hamyon_configured():
    return bool(HAMYON_SHOP_ID and HAMYON_SHOP_KEY)


def hamyon_call(method, path, params, timeout=15):
    """HamyonAPI ga so'rov. (http_kod, json_dict) qaytaradi."""
    qs = urllib.parse.urlencode(params)
    if method == "GET":
        req = urllib.request.Request(HAMYON_BASE + path + "?" + qs)
    else:
        req = urllib.request.Request(HAMYON_BASE + path, data=qs.encode(), method="POST",
                                     headers={"Content-Type": "application/x-www-form-urlencoded"})
    try:
        with urllib.request.urlopen(req, timeout=timeout) as r:
            code, body = r.status, r.read().decode("utf-8", "replace")
    except urllib.error.HTTPError as e:
        code, body = e.code, e.read().decode("utf-8", "replace")
    except Exception:
        raise PayError("To'lov tizimiga ulanib bo'lmadi. Birozdan keyin qayta urinib ko'ring.")
    try:
        data = json.loads(body)
    except ValueError:
        data = {"error": body[:200]}
    return code, data if isinstance(data, dict) else {"data": data}


def create_remote_payment(tid, months):
    """To'lov yaratadi. Bir xil summada ochiq to'lov bo'lsa, summani 1 so'mdan oshirib qayta urinadi
    (HamyonAPI qoidasi: bitta do'konda bir vaqtda bir xil summali ikkita ochiq to'lov bo'lmaydi)."""
    base = PRICE_MONTH * months
    order_id = "SUB-%d-%s" % (tid, secrets.token_hex(4))
    for off in range(10):
        amount = base + off
        code, j = hamyon_call("POST", "/payment/create", dict(
            shop_id=HAMYON_SHOP_ID, shop_key=HAMYON_SHOP_KEY, amount=amount, order_id=order_id))
        if code == 200 and j.get("payment_id"):
            try:
                if int(float(j.get("amount", amount))) != amount:
                    raise PayError("To'lov summasi mos kelmadi. Qayta urinib ko'ring.")
            except (TypeError, ValueError):
                pass
            return amount, order_id, j
        err = str(j.get("error") or j.get("message") or "")
        if "ochiq" in err.lower() and "summa" in err.lower():
            continue
        raise PayError(err or "To'lov yaratib bo'lmadi.")
    raise PayError("Hozir to'lov qabul qilish band. Bir necha daqiqadan keyin qayta urinib ko'ring.")


def _expire_dt(j):
    try:
        return datetime.fromtimestamp(int(j.get("expire_at")), TZ).replace(tzinfo=None)
    except (TypeError, ValueError):
        pass
    try:
        return now() + timedelta(seconds=int(j.get("expires_in")))
    except (TypeError, ValueError):
        return now() + timedelta(minutes=5)


def mark_paid(pid, raw=None):
    """To'lovni 'paid' qiladi va obunani uzaytiradi. Bir to'lov faqat bir marta hisoblanadi."""
    conn = db()
    cur = conn.execute("UPDATE payments SET status='paid', paid_at=?, raw=? WHERE id=? AND status!='paid'",
                       (iso(now()), (raw or "")[:2000], pid))
    if cur.rowcount == 0:
        conn.commit()
        return False
    p = conn.execute("SELECT * FROM payments WHERE id=?", (pid,)).fetchone()
    extend_subscription(p["teacher_id"], p["months"])
    conn.commit()
    return True


def pending_payment(tid):
    ex("UPDATE payments SET status='expired' WHERE teacher_id=? AND status='pending' AND expire_at < ?",
       (tid, iso(now())))
    return q("SELECT * FROM payments WHERE teacher_id=? AND status='pending' ORDER BY id DESC LIMIT 1", (tid,), one=True)


@app.route("/teacher/billing")
@role_required("teacher")
def t_billing():
    uid = current_user()["id"]
    pend = pending_payment(uid)
    left = max(0, int((parse(pend["expire_at"]) - now()).total_seconds())) if pend else 0
    history = q("SELECT * FROM payments WHERE teacher_id=? ORDER BY id DESC LIMIT 20", (uid,))
    return render_template("t_billing.html", pend=pend, left=left, history=history, months=MONTH_CHOICES,
                           configured=hamyon_configured(), trial_days=TRIAL_DAYS)


@app.route("/teacher/billing/pay", methods=["POST"])
@role_required("teacher")
def t_billing_pay():
    uid = current_user()["id"]
    if not hamyon_configured():
        flash("To'lov tizimi hali sozlanmagan. Administratorga murojaat qiling.", "err")
        return redirect(url_for("t_billing"))
    months = as_int(request.form.get("months"), 1, 1, 12)
    if months not in MONTH_CHOICES:
        months = 1
    if pending_payment(uid):
        flash("Sizda tugallanmagan to'lov bor. Avval uni yakunlang yoki bekor qiling.", "err")
        return redirect(url_for("t_billing"))
    try:
        amount, order_id, j = create_remote_payment(uid, months)
    except PayError as e:
        flash("To'lov yaratilmadi: %s" % e, "err")
        return redirect(url_for("t_billing"))
    ex("INSERT INTO payments(teacher_id,order_id,payment_id,months,amount,card,status,created_at,expire_at) "
       "VALUES(?,?,?,?,?,?,'pending',?,?)",
       (uid, order_id, str(j["payment_id"]), months, amount, str(j.get("card", "")), iso(now()), iso(_expire_dt(j))))
    return redirect(url_for("t_billing"))


@app.route("/teacher/billing/cancel", methods=["POST"])
@role_required("teacher")
def t_billing_cancel():
    uid = current_user()["id"]
    p = pending_payment(uid)
    if p:
        try:
            hamyon_call("POST", "/payment/cancel", dict(shop_id=HAMYON_SHOP_ID, shop_key=HAMYON_SHOP_KEY,
                                                        payment_id=p["payment_id"]))
        except PayError:
            pass
        ex("UPDATE payments SET status='cancelled' WHERE id=? AND status='pending'", (p["id"],))
        flash("To'lov bekor qilindi.", "ok")
    return redirect(url_for("t_billing"))


@app.route("/teacher/billing/status/<int:pid>")
@role_required("teacher")
def t_billing_status(pid):
    p = q("SELECT status FROM payments WHERE id=? AND teacher_id=?", (pid, current_user()["id"]), one=True)
    if not p:
        abort(404)
    return jsonify(status=p["status"])


@app.route("/teacher/billing/check", methods=["POST"])
@role_required("teacher")
def t_billing_check():
    """Callback kelmay qolgan holatda to'lovni HamyonAPI dan qo'lda so'rab tekshirish."""
    uid = current_user()["id"]
    p = q("SELECT * FROM payments WHERE teacher_id=? AND status IN ('pending','expired') AND created_at > ? "
          "ORDER BY id DESC LIMIT 1", (uid, iso(now() - timedelta(hours=2))), one=True)
    if not p or not p["payment_id"]:
        flash("Tekshiriladigan to'lov topilmadi.", "err")
        return redirect(url_for("t_billing"))
    try:
        code, j = hamyon_call("GET", "/payment/status", dict(shop_id=HAMYON_SHOP_ID, shop_key=HAMYON_SHOP_KEY,
                                                             payment_id=p["payment_id"]))
    except PayError as e:
        flash(str(e), "err")
        return redirect(url_for("t_billing"))
    inner = j.get("payment") if isinstance(j.get("payment"), dict) else {}
    st = str(j.get("status") or inner.get("status") or "").lower()
    amt = j.get("amount", inner.get("amount"))
    amount_ok = True
    if amt is not None:
        try:
            amount_ok = int(float(amt)) == p["amount"]
        except (TypeError, ValueError):
            amount_ok = False
    if st == "paid" and amount_ok:
        if mark_paid(p["id"], json.dumps(j, ensure_ascii=False)):
            flash("To'lov qabul qilindi. Obuna uzaytirildi!", "ok")
    else:
        flash("To'lov hali tushmagan.", "err")
    return redirect(url_for("t_billing"))


# ---- HamyonAPI callback'lari (prepare_url / complete_url) ----
PAID_WORDS = {"paid", "success", "completed", "done"}
CANCEL_WORDS = {"cancel", "cancelled", "canceled", "expired", "failed"}


def hook_values():
    v = {}
    v.update(request.args.to_dict())
    v.update(request.form.to_dict())
    j = request.get_json(silent=True)
    if isinstance(j, dict):
        for k, x in j.items():
            v[k] = x if isinstance(x, str) else (json.dumps(x) if isinstance(x, (dict, list)) else str(x))
    return v


def verify_sign(vals):
    """sign = md5(shop_id + payment_id + amount + shop_key)"""
    sign = str(vals.get("sign", "")).strip().lower()
    if not sign or not HAMYON_SHOP_KEY:
        return False
    shop_id, pid, amount = str(vals.get("shop_id", "")), str(vals.get("payment_id", "")), str(vals.get("amount", ""))
    cands = {amount}
    try:
        cands.add(str(int(float(amount))))
    except ValueError:
        pass
    for a in cands:
        exp = hashlib.md5((shop_id + pid + a + HAMYON_SHOP_KEY).encode()).hexdigest()
        if hmac.compare_digest(exp, sign):
            return True
    return False


def handle_hook(kind):
    vals = hook_values()
    safe = {k: v for k, v in vals.items() if k.lower() not in ("shop_key", "key", "secret")}

    def done(resp, code, result):
        ex("INSERT INTO hook_log(at,path,ip,body,result) VALUES(?,?,?,?,?)",
           (iso(now()), request.path, request.remote_addr, json.dumps(safe, ensure_ascii=False)[:2000], result))
        return jsonify(resp), code

    if not hamyon_configured():
        return done({"status": "error", "error": "not configured"}, 503, "not configured")
    if str(vals.get("shop_id", "")) != HAMYON_SHOP_ID or not verify_sign(vals):
        return done({"status": "error", "error": "bad signature"}, 403, "bad signature")
    pid, oid = str(vals.get("payment_id", "")), str(vals.get("order_id", ""))
    p = q("SELECT * FROM payments WHERE payment_id=?", (pid,), one=True)
    if not p and oid:
        p = q("SELECT * FROM payments WHERE order_id=?", (oid,), one=True)
    if not p:
        return done({"status": "error", "error": "payment not found"}, 404, "not found")
    try:
        amt = int(float(vals.get("amount")))
    except (TypeError, ValueError):
        amt = -1
    if amt != p["amount"]:
        return done({"status": "error", "error": "amount mismatch"}, 400, "amount mismatch")
    status = str(vals.get("status", "")).strip().lower()
    if kind == "auto":
        kind = "complete" if status else "prepare"
    if kind == "prepare":
        ex("UPDATE payments SET prepared_at=? WHERE id=?", (iso(now()), p["id"]))
        return done({"status": "ok"}, 200, "prepared")
    if status in PAID_WORDS:
        credited = mark_paid(p["id"], json.dumps(safe, ensure_ascii=False))
        return done({"status": "ok"}, 200, "paid (credited)" if credited else "paid (already)")
    if status in CANCEL_WORDS:
        ex("UPDATE payments SET status='cancelled' WHERE id=? AND status='pending'", (p["id"],))
        return done({"status": "ok"}, 200, "cancelled")
    return done({"status": "error", "error": "unknown status"}, 400, "unknown status: %s" % status[:30])


@app.route("/hooks/hamyon", methods=["GET", "POST"])
def hook_auto():
    return handle_hook("auto")


@app.route("/hooks/hamyon/prepare", methods=["GET", "POST"])
def hook_prepare():
    return handle_hook("prepare")


@app.route("/hooks/hamyon/complete", methods=["GET", "POST"])
def hook_complete():
    return handle_hook("complete")


# --------------------------------------------------------------------------
# Rasmlar (media), xavfsizlik sarlavhalari, diagnostika
# --------------------------------------------------------------------------
@app.route("/media/<name>")
def media(name):
    """Savol/variant rasmlari. Faqat tizimga kirgan foydalanuvchilarga; fayl nomlari tasodifiy (taxmin qilib bo'lmaydi)."""
    if not current_user() or not MEDIA_RE.match(name):
        abort(404)
    path = os.path.join(UPLOAD_DIR, name)
    if not os.path.isfile(path):
        abort(404)
    resp = send_file(path, mimetype="image/png" if name.endswith(".png") else "image/jpeg")
    resp.headers["Cache-Control"] = "private, max-age=86400"
    return resp


@app.after_request
def secure_headers(resp):
    resp.headers.setdefault("X-Content-Type-Options", "nosniff")
    resp.headers.setdefault("X-Frame-Options", "SAMEORIGIN")
    resp.headers.setdefault("Referrer-Policy", "same-origin")
    if resp.mimetype in ("text/html", "application/json"):
        resp.headers["Cache-Control"] = "no-store"  # hosting/CDN sahifalarni keshlab qo'ymasin (kirish "yo'qolishi" sababi bo'lishi mumkin)
    return resp


@app.route("/healthz")
def healthz():
    """Server sozlamalarini tekshirish. Sahifani bir necha marta yangilang: 'secret_fp' har doim BIR XIL
    bo'lishi kerak (turli bo'lsa — workerlar turli sessiya kaliti ishlatyapti)."""
    ok_db = True
    try:
        q("SELECT 1", one=True)
    except Exception:
        ok_db = False
    return jsonify(
        ok=ok_db, pid=os.getpid(), secret_source=SECRET_SOURCE,
        secret_fp=hashlib.sha256(SECRET.encode()).hexdigest()[:8],
        scheme=request.scheme, cookie_secure=app.config["SESSION_COOKIE_SECURE"],
        proxy_headers=bool(request.headers.get("X-Forwarded-For") or request.headers.get("X-Forwarded-Proto")),
        db_writable=os.access(DB_PATH, os.W_OK), uploads_writable=os.access(UPLOAD_DIR, os.W_OK),
        pillow=Image is not None, tz=str(TZ), hamyon_configured=hamyon_configured())


init_db()
print("[i] sessiya kaliti: %s | baza: %s | rasmlar: %s" % (SECRET_SOURCE, DB_PATH, UPLOAD_DIR))

if __name__ == "__main__":
    host = os.environ.get("HOST", "0.0.0.0")
    port = int(os.environ.get("PORT", "5000"))
    app.run(host=host, port=port, debug=False, threaded=True)
